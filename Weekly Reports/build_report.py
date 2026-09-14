import copy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PROJECT = "Self Driving RC Car"
DATE = "9/15/2026"
SUMMARY_DATE = "9/15/2026"

members = ["Ethan Wells (Group Leader)", "Alexis Perez", "Ethan Bishop", "Abigail Duran"]

summary = ("This week the team made progress on all fronts. "
    "Ethan Wells trained the traffic-sign model — a convolutional neural network on the GTSRB German dataset reaching about 95% validation accuracy — "
    "committed the model and training code, and ordered the CR2032 RTC battery and the KIOXIA NVMe boot SSD with a USB enclosure. "
    "Alexis Perez bench-tested the ultrasonic sensor's distance readings and accuracy and researched drive-motor options. "
    "Ethan Bishop tested the IMU for measurement accuracy, researched motor-controller options, and started drying filament to prepare the home 3D printer for chassis-component testing. "
    "Abigail Duran completed and refined the first-semester schematic around the Raspberry Pi Compute Module 5, resolved ERC issues to ready it for PCB development, "
    "and began hardware prototyping for the stationary test setup.")

plan = [
    ("Ethan Wells", "Prepare the LISA dataset (US traffic signs) for training; test the trained German-sign (GTSRB) model."),
    ("Alexis Perez", "Select and order motors; write code to integrate the ultrasonic sensor on the Raspberry Pi."),
    ("Ethan Bishop", "Select and order a motor controller; write code to integrate the IMU on the Raspberry Pi."),
    ("Abigail Duran", "Begin the PCB design and layout from the completed schematic; place major components and determine board dimensions; begin planning power and signal routing."),
]

contributions = [
    ("Ethan Wells", [("9/10/2026", "Built and ran the first CNN training on the GTSRB German traffic-sign dataset in TensorFlow", 3),
                        ("9/14/2026", "Tuned the model to ~95% validation accuracy and committed the trained model (gtsrb_model.keras) and training script to the repo", 4)]),
    ("Alexis Perez", [("", "Tested the ultrasonic sensor's distance readings and measurement accuracy", None),
                        ("", "Researched drive-motor options for the vehicle", None)]),
    ("Ethan Bishop", [("", "Tested the IMU for measurement accuracy", None),
                        ("", "Researched motor-controller options", None),
                        ("", "Started drying filament to prep the home 3D printer for chassis-component testing", None)]),
    ("Abigail Duran", [("", "Completed the first-semester schematic for stationary testing of the autonomous car system", None),
                        ("", "Built and refined the schematic around the Raspberry Pi Compute Module 5", None),
                        ("", "Resolved schematic/ERC issues and prepared the design for PCB development", None),
                        ("", "Began hardware prototyping for the stationary test setup", None)]),
]

hours = [("Ethan Wells", "7", "12"), ("Alexis Perez", "6", "9"),
         ("Ethan Bishop", "6", "9"), ("Abigail Duran", "6", "11")]

d = Document("project_name.docx")

def para_by(pred):
    for p in d.paragraphs:
        if pred(p.text):
            return p
    raise SystemExit("anchor not found: " + repr(pred))

def add_run(p_el, text, bold=False, underline=False):
    r = OxmlElement('w:r'); rpr = OxmlElement('w:rPr')
    if bold:
        rpr.append(OxmlElement('w:b'))
    if underline:
        u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rpr.append(u)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '24'); rpr.append(sz)   # 12pt
    r.append(rpr)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text; r.append(t)
    p_el.append(r)
    return r

# normal-body paragraph template (clone the summary instruction paragraph, keep its pPr)
p_sum_head = para_by(lambda t: t.startswith("Weekly Summary"))
sum_instr = p_sum_head._p.getnext()          # the instruction paragraph after the heading
normal_tmpl = copy.deepcopy(sum_instr)

def new_normal():
    el = copy.deepcopy(normal_tmpl)
    for r in el.findall(qn('w:r')): el.remove(r)
    return el

def line_para(runs):
    """runs = list of (text, bold, underline)"""
    el = new_normal()
    for text, b, u in runs:
        add_run(el, text, b, u)
    return el

# 1) Project Title
p_title = para_by(lambda t: t.startswith("Project Title"))
rs = p_title.runs
rs[0].text = "Project Title: "
rs[1].text = PROJECT
for r in rs[2:]: r.text = ""

# 2) Date
p_date = para_by(lambda t: t.startswith("Date:"))
rs = p_date.runs
rs[0].text = "Date: "
rs[1].text = DATE
for r in rs[2:]: r.text = ""

# 3) Project Members -> insert one line per member after the heading
p_mem = para_by(lambda t: t.strip() == "Project Members:")
anchor = p_mem._p
# remove the single blank paragraph the template puts right after members, if present
nxt = anchor.getnext()
if nxt is not None and nxt.tag == qn('w:p') and not nxt.findall(qn('w:r')):
    nxt.getparent().remove(nxt)
for m in members:
    el = line_para([(m, False, False)])
    anchor.addnext(el); anchor = el

# 4) Weekly Summary heading date + body
sum_runs = p_sum_head.runs   # capture once: p_sum_head.runs[0] would be a fresh object each call
for i, r in enumerate(sum_runs):
    r.text = "Weekly Summary (%s):" % SUMMARY_DATE if i == 0 else ""
# replace instruction paragraph content with the summary
for r in list(sum_instr.findall(qn('w:r'))): sum_instr.remove(r)
add_run(sum_instr, summary, False, False)

# 5) Proposed Plan -> per-member lines replacing the instruction paragraph
p_plan_head = para_by(lambda t: t.startswith("Proposed Plan"))
plan_instr = p_plan_head._p.getnext()
# turn the instruction paragraph into the first member's line
for r in list(plan_instr.findall(qn('w:r'))): plan_instr.remove(r)
n0, g0 = plan[0]
add_run(plan_instr, n0 + ": ", False, True)
add_run(plan_instr, g0, False, False)
anchor = plan_instr
for n, g in plan[1:]:
    el = line_para([(n + ": ", False, True), (g, False, False)])
    anchor.addnext(el); anchor = el

# 6) Weekly Contributions -> remove everything between the heading and Hour Tracker, rebuild
p_contrib = para_by(lambda t: t.strip() == "Weekly Contributions:")
p_hour = para_by(lambda t: t.strip() == "Hour Tracker:")
start = p_contrib._p
stop = p_hour._p
# collect elements strictly between start and stop
to_remove = []
cur = start.getnext()
while cur is not None and cur is not stop:
    to_remove.append(cur)
    cur = cur.getnext()
for el in to_remove:
    el.getparent().remove(el)
# rebuild contributions after the heading
anchor = start
for name, entries in contributions:
    head = line_para([(name + ":", False, True)])
    anchor.addnext(head); anchor = head
    for date, text, hrs in entries:
        line = ("%s – %s" % (date, text)) if date else text   # date optional
        if hrs is not None:                                   # hours optional (placeholder)
            unit = "hour" if hrs == 1 else "hours"
            line += " (%d %s)" % (hrs, unit)
        el = line_para([(line, False, False)])
        anchor.addnext(el); anchor = el
# one blank spacer before Hour Tracker
spacer = new_normal()
anchor.addnext(spacer)

# 7) Hour Tracker table
tbl = d.tables[0]
def set_cell(cell, text):
    p = cell.paragraphs[0]
    for r in list(p._p.findall(qn('w:r'))): p._p.remove(r)
    add_run(p._p, text, False, False)
for i, (name, wk, cum) in enumerate(hours):
    row = tbl.rows[i + 1]
    set_cell(row.cells[0], name)
    set_cell(row.cells[1], wk)
    set_cell(row.cells[2], cum)

d.save("Self_Driving_RC_Car.docx")
print("saved")
